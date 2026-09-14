"""
ZFJ HTML Parser + Timesheet Filler — Vercel Serverless
Entry point for Vercel Python runtime.
"""

import os
import io
import time
from datetime import datetime
from collections import defaultdict

from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import lxml.html
import openpyxl
import openpyxl.styles

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

# ============================================================
# Constants
# ============================================================
COL_DATE = 1
COL_START = 2
COL_END = 3
COL_PRESENT = 5
COL_ACTIVITY = 11
COL_PROJECT_NAME = 12
COL_PROJECT_ID = 13
COL_APLIKASI = 14
COL_AIP_FITUR = 15
COL_DIVISI = 16
COL_DEPARTEMENT = 17
COL_SUB_DEPT = 18

ROW_PROJECT_NAME = 1
ROW_UNIT_DIVISI = 2
ROW_NAME = 3
ROW_MII_ID = 4
ROW_SITE = 5

DATA_ROW_START = 9

# ============================================================
# Date Helpers
# ============================================================
MONTH_MAP = {
    "Jan": "01", "Feb": "02", "Mar": "03", "Apr": "04",
    "May": "05", "Jun": "06", "Jul": "07", "Aug": "08",
    "Sep": "09", "Oct": "10", "Nov": "11", "Dec": "12",
}


def convert_date(raw: str) -> str:
    try:
        parts = raw.strip().split("/")
        if len(parts) != 3:
            return raw
        day = parts[0].zfill(2)
        month = MONTH_MAP.get(parts[1], parts[1])
        year = parts[2]
        if len(year) == 2:
            year = f"20{year}"
        return f"{day}-{month}-{year}"
    except Exception:
        return raw


def sort_date_key(date_str: str):
    try:
        return datetime.strptime(date_str, "%d-%m-%Y")
    except ValueError:
        return datetime.min


def is_holiday_row(ws, row: int) -> bool:
    fill = ws.cell(row, COL_DATE).fill
    if fill.patternType == "solid":
        try:
            if fill.fgColor.theme == 0 and fill.fgColor.tint < 0:
                return True
        except (AttributeError, TypeError):
            pass
    return False


# ============================================================
# ZFJ Parser
# ============================================================
def parse_zfj_html(html_content: str) -> list[dict]:
    tree = lxml.html.fromstring(html_content)
    grouped = defaultdict(list)

    for tr in tree.iter("tr"):
        tds = tr.findall("td")
        if len(tds) < 20:
            continue

        execution_id = (tds[0].text_content() or "").strip()
        executed_on = (tds[11].text_content() or "").strip()

        if not execution_id or not executed_on:
            continue

        tanggal = convert_date(executed_on)
        tc_key = f"MAV-{execution_id}"

        if tc_key not in grouped[tanggal]:
            grouped[tanggal].append(tc_key)

    result = []
    for tanggal in sorted(grouped.keys(), key=sort_date_key):
        result.append({"tanggal": tanggal, "tc": sorted(grouped[tanggal])})

    return result


# ============================================================
# Timesheet Filler
# ============================================================
def fill_timesheet(
    excel_bytes, parsed_data,
    name_project="", unit_divisi="", name="",
    mii_id="", site="", project_name="MMP", sign_date="",
):
    from openpyxl.styles import Alignment as Al

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    # Header
    if name_project:
        ws.cell(ROW_PROJECT_NAME, 3).value = f": {name_project}"
    if unit_divisi:
        ws.cell(ROW_UNIT_DIVISI, 3).value = f": {unit_divisi}"
    if name:
        ws.cell(ROW_NAME, 3).value = f": {name}"
    if mii_id:
        ws.cell(ROW_MII_ID, 3).value = f": {mii_id}"
    if site:
        ws.cell(ROW_SITE, 3).value = f": {site}"

    # Date lookup
    date_lookup = {}
    for item in parsed_data:
        try:
            dt = datetime.strptime(item["tanggal"], "%d-%m-%Y")
            date_lookup[dt.date()] = item["tc"]
        except ValueError:
            continue

    stats = {"kerja_filled": 0, "libur_filled": 0, "skipped": 0}

    for row in range(DATA_ROW_START, DATA_ROW_START + 31):
        cell_date = ws.cell(row, COL_DATE).value
        if cell_date is None:
            continue
        if isinstance(cell_date, datetime):
            row_date = cell_date.date()
        else:
            continue

        if row_date not in date_lookup:
            stats["skipped"] += 1
            continue

        tc_list = date_lookup[row_date]
        tc_text = ", ".join(tc_list)
        holiday = is_holiday_row(ws, row)

        if holiday:
            existing = ws.cell(row, COL_ACTIVITY).value
            if existing and str(existing).strip():
                ws.cell(row, COL_ACTIVITY).value = f"{existing}\n{tc_text}"
            else:
                ws.cell(row, COL_ACTIVITY).value = tc_text
            stats["libur_filled"] += 1
        else:
            existing = ws.cell(row, COL_ACTIVITY).value
            if existing and str(existing).strip():
                ws.cell(row, COL_ACTIVITY).value = f"{existing}\n{tc_text}"
            else:
                ws.cell(row, COL_ACTIVITY).value = tc_text

            ws.cell(row, COL_PRESENT).value = "P"

            center_align = openpyxl.styles.Alignment(horizontal="center", vertical="center")
            ws.cell(row, COL_PROJECT_NAME).value = project_name
            ws.cell(row, COL_PROJECT_NAME).alignment = center_align
            ws.cell(row, COL_PROJECT_ID).value = "P24005"
            ws.cell(row, COL_PROJECT_ID).alignment = center_align
            ws.cell(row, COL_APLIKASI).value = "Wondr"
            ws.cell(row, COL_APLIKASI).alignment = center_align
            ws.cell(row, COL_AIP_FITUR).value = "-"
            ws.cell(row, COL_AIP_FITUR).alignment = center_align
            ws.cell(row, COL_DIVISI).value = "RDL"
            ws.cell(row, COL_DIVISI).alignment = center_align
            ws.cell(row, COL_DEPARTEMENT).value = "Maverick"
            ws.cell(row, COL_DEPARTEMENT).alignment = center_align
            ws.cell(row, COL_SUB_DEPT).value = "-"
            ws.cell(row, COL_SUB_DEPT).alignment = center_align

            stats["kerja_filled"] += 1

    # Signature
    if name:
        ws.cell(43, 1).value = f"(  {name}  )"
    ws.cell(43, 4).value = "(  Saiful Fahmi  )"
    ws.cell(43, 7).value = "(  Yohanes Hendra Jasin  )"

    ws.cell(43, 1).alignment = Al(horizontal="center", vertical="bottom", wrap_text=True)
    ws.cell(43, 4).alignment = Al(horizontal="center", vertical="bottom", wrap_text=True)
    ws.cell(43, 7).alignment = Al(horizontal="center", vertical="bottom", wrap_text=True)

    for c in [1, 4, 7]:
        existing_date = ws.cell(47, c).value
        ws.cell(48, c).value = None

    ws.cell(47, 1).value = "Karyawan"
    ws.cell(47, 4).value = "Chapter Lead"
    ws.cell(47, 7).value = "Departement Head"

    for merge_range in ["A47:C47", "D47:F47", "G47:J47"]:
        ws.merge_cells(merge_range)

    for c in [1, 4, 7]:
        ws.cell(47, c).alignment = Al(horizontal="center", vertical="center")
        ws.cell(47, c).font = openpyxl.styles.Font(name="Arial", size=10, bold=True)

    date_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style="thin"),
        right=openpyxl.styles.Side(style="thin"),
        top=openpyxl.styles.Side(style="thin"),
        bottom=openpyxl.styles.Side(style="thin"),
    )
    if sign_date:
        ws.cell(48, 1).value = f"DATE : {sign_date}"
        ws.cell(48, 4).value = f"DATE : {sign_date}"
        ws.cell(48, 7).value = f"DATE : {sign_date}"
    else:
        ws.cell(48, 1).value = "DATE :"
        ws.cell(48, 4).value = "DATE :"
        ws.cell(48, 7).value = "DATE :"

    for merge_range in ["A48:C48", "D48:F48", "G48:J48"]:
        ws.merge_cells(merge_range)
    for c in range(1, 11):
        ws.cell(48, c).border = date_border

    ws.print_area = "A1:R48"

    # Auto-fit
    COL_K_WIDTH_CHARS = 80
    LINE_HEIGHT = 15
    for row in range(DATA_ROW_START, DATA_ROW_START + 31):
        activity = ws.cell(row, COL_ACTIVITY).value
        if activity and len(activity) > COL_K_WIDTH_CHARS:
            num_lines = (len(activity) // COL_K_WIDTH_CHARS) + 1
            ws.row_dimensions[row].height = max(LINE_HEIGHT * num_lines, 15)
            ws.cell(row, COL_ACTIVITY).alignment = openpyxl.styles.Alignment(
                wrap_text=True, vertical="top"
            )
    ws.column_dimensions["K"].width = 60

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), stats


# ============================================================
# Routes
# ============================================================
@app.route("/", methods=["GET"])
def index():
    return jsonify({
        "service": "ZFJ HTML Parser + Timesheet Filler",
        "version": "4.1.0-vercel",
        "endpoints": {
            "GET /": "Health check",
            "POST /parse": "Upload ZFJ HTML → JSON",
            "POST /fill-preview": "Preview filled data (JSON)",
            "POST /fill": "Download filled Excel",
        },
    })


def _get_form_fields():
    return {
        "name_project": request.form.get("name_project", ""),
        "unit_divisi": request.form.get("unit_divisi", ""),
        "name": request.form.get("name", ""),
        "mii_id": request.form.get("mii_id", ""),
        "site": request.form.get("site", ""),
        "project_name": request.form.get("project_name", "MMP"),
        "sign_date": request.form.get("sign_date", ""),
    }


def _validate_files():
    if "html" not in request.files:
        return None, None, ("Missing 'html' field.", 400)
    if "excel" not in request.files:
        return None, None, ("Missing 'excel' field.", 400)

    html_file = request.files["html"]
    excel_file = request.files["excel"]

    if not html_file.filename.lower().endswith((".html", ".htm")):
        return None, None, ("HTML file must be .html/.htm", 400)
    if not excel_file.filename.lower().endswith((".xlsx",)):
        return None, None, ("Excel file must be .xlsx", 400)

    raw_html = html_file.read()
    try:
        html_content = raw_html.decode("utf-8")
    except UnicodeDecodeError:
        html_content = raw_html.decode("latin-1")

    excel_bytes = excel_file.read()
    return html_content, excel_bytes, None


@app.route("/parse", methods=["POST"])
def parse():
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file."}), 400
    file = request.files["file"]
    raw = file.read()
    try:
        content = raw.decode("utf-8")
    except UnicodeDecodeError:
        content = raw.decode("latin-1")
    start = time.time()
    data = parse_zfj_html(content)
    elapsed = round(time.time() - start, 3)
    total = sum(len(i["tc"]) for i in data)
    return jsonify({"status": "success", "total_dates": len(data), "total_executions": total, "parse_time_seconds": elapsed, "data": data})


@app.route("/fill-preview", methods=["POST"])
def fill_preview():
    html_content, excel_bytes, error = _validate_files()
    if error:
        return jsonify({"status": "error", "message": error[0]}), error[1]
    fields = _get_form_fields()
    start = time.time()
    parsed_data = parse_zfj_html(html_content)
    filled_bytes, stats = fill_timesheet(excel_bytes, parsed_data, **fields)
    elapsed = round(time.time() - start, 3)

    wb = openpyxl.load_workbook(io.BytesIO(filled_bytes))
    ws = wb.active
    preview = []
    for row in range(DATA_ROW_START, DATA_ROW_START + 31):
        cell_date = ws.cell(row, COL_DATE).value
        if cell_date is None:
            continue
        activity = ws.cell(row, COL_ACTIVITY).value
        present = ws.cell(row, COL_PRESENT).value
        date_str = cell_date.strftime("%d-%m-%Y") if isinstance(cell_date, datetime) else str(cell_date)
        preview.append({
            "tanggal": date_str,
            "present": present or "",
            "activity": (activity[:80] + "...") if activity and len(activity) > 80 else (activity or ""),
            "tc_count": len(activity.split(", ")) if activity else 0,
        })

    return jsonify({"status": "success", "stats": stats, "parse_time_seconds": elapsed, "preview": preview})


@app.route("/fill", methods=["POST"])
def fill():
    html_content, excel_bytes, error = _validate_files()
    if error:
        return jsonify({"status": "error", "message": error[0]}), error[1]
    fields = _get_form_fields()
    parsed_data = parse_zfj_html(html_content)
    filled_bytes, stats = fill_timesheet(excel_bytes, parsed_data, **fields)
    original_name = request.files["excel"].filename.rsplit(".", 1)[0]
    return send_file(
        io.BytesIO(filled_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"{original_name}_FILLED.xlsx",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=4000, debug=True)